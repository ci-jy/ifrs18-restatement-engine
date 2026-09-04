import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from ifrs18.cli import main

ROOT = Path(__file__).resolve().parents[1]


def test_restate_and_verify_cli(tmp_path, sample_csv, capsys):
    out = tmp_path / "restated.xlsx"
    assert main(["restate", str(sample_csv), "--entity-type", "general", "-o", str(out)]) == 0
    text = capsys.readouterr().out
    assert "Net income conserved" in text and "Adjusted EBITDA FY2026: 343,018.85" in text
    assert set(load_workbook(out).sheetnames) >= {"Original", "Mapping", "Restated", "Checks", "MPM"}
    assert main(["verify-workbook", str(out), "--tolerance", "1e-9"]) == 0


def test_investing_entity_moves_returns_into_operating(tmp_path, sample_csv, capsys):
    out = tmp_path / "inv.xlsx"
    main(["restate", str(sample_csv), "--entity-type", "investing", "-o", str(out)])
    ws = load_workbook(out)["Mapping"]
    cats = {r[0].value: r[2].value for r in ws.iter_rows(min_row=2)}
    assert cats["L011"] == cats["L012"] == cats["L013"] == "operating"
    assert cats["L009"] == "investing"  # equity-accounted stays investing


def test_mapping_override(tmp_path, sample_csv):
    ov = tmp_path / "ov.csv"
    ov.write_text("line_id,category\nL015,financing\n")
    out = tmp_path / "o.xlsx"
    assert main(["restate", str(sample_csv), "--mapping", str(ov), "-o", str(out)]) == 0
    cats = {r[0].value: r[2].value for r in load_workbook(out)["Mapping"].iter_rows(min_row=2)}
    assert cats["L015"] == "financing"


def test_module_entry_point(tmp_path, sample_csv):
    out = tmp_path / "m.xlsx"
    p = subprocess.run([sys.executable, "-m", "ifrs18", "restate", str(sample_csv), "-o", str(out)], cwd=ROOT,
                       capture_output=True, text=True)
    assert p.returncode == 0, p.stderr
