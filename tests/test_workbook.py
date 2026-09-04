"""Excel round trip: LibreOffice recalculates the workbook and must match Python within 1e-9."""
import random
from decimal import Decimal

from openpyxl import load_workbook

from ifrs18.mpm import load_definition, parse_definition
from ifrs18.restate import restate
from ifrs18.rules import CATEGORIES, classify
from ifrs18.statement import ITEM, PROFIT, Line, Statement, read_csv
from ifrs18.workbook import verify_workbook, write_workbook


def test_sample_workbook_recalculates_in_libreoffice_to_1e9(tmp_path, sample_csv, sample_mpm):
    stmt = read_csv(sample_csv)
    classes = classify(stmt)
    r = restate(stmt, {k: c.category for k, c in classes.items()})
    out = tmp_path / "restated.xlsx"
    expected = write_workbook(r, classes, out, mpm=load_definition(sample_mpm))
    assert any(sheet == "MPM" for sheet, _ in expected)
    res = verify_workbook(out, 1e-9)
    assert res.ok, res.failures
    assert res.checked == len(expected) and res.max_abs_diff <= 1e-9


def test_workbook_cells_are_formulas_not_pasted_values(tmp_path, sample_csv, sample_mpm):
    stmt = read_csv(sample_csv)
    classes = classify(stmt)
    r = restate(stmt, {k: c.category for k, c in classes.items()})
    out = tmp_path / "w.xlsx"
    expected = write_workbook(r, classes, out, mpm=load_definition(sample_mpm))
    wb = load_workbook(out)
    for sheet, cell in expected:
        v = wb[sheet][cell].value
        assert isinstance(v, str) and v.startswith("="), (sheet, cell, v)


def test_changing_a_category_in_excel_flows_through(tmp_path, sample_csv):
    """Tamper with the reviewer input: recalculated totals move, profit does not."""
    stmt = read_csv(sample_csv)
    classes = classify(stmt)
    r = restate(stmt, {k: c.category for k, c in classes.items()})
    out = tmp_path / "w.xlsx"
    write_workbook(r, classes, out)
    wb = load_workbook(out)
    ws = wb["Mapping"]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == "L009":          # associate: investing -> operating
            row[2].value = "operating"
    wb.save(out)
    res = verify_workbook(out, 1e-9)
    bad = {(s, c) for s, c, _, _ in res.failures}
    assert bad, "a changed category must change the recalculated category totals"
    assert all(s == "Restated" for s, _ in bad)          # Checks sheet stays at zero
    profit_rows = [c for s, c, _, _ in res.failures if s == "Checks"]
    assert not profit_rows


def test_random_statement_with_mpm_round_trips(tmp_path):
    rnd = random.Random(7)
    periods = ["FY2026", "FY2025"]
    lines = [Line(f"L{i:02d}", f"Line {i}", {p: Decimal(rnd.randint(-10**9, 10**9)) / 100 for p in periods}) for i in range(25)]
    lines.append(Line("P", "Profit", {p: sum((l.amounts[p] for l in lines), Decimal(0)) for p in periods}, kind=PROFIT))
    stmt = Statement("Random plc", periods, lines)
    mapping = {l.line_id: rnd.choice(CATEGORIES) for l in lines if l.kind == ITEM}
    r = restate(stmt, mapping)
    mpm = parse_definition({"name": "Adjusted EBIT", "subtotal": "pbfit", "adjustments": [
        {"label": "Line 3 add-back", "line": "L03", "tax_rate": "0.2537", "nci_share": "0.1234"},
        {"label": "Note amount", "amounts": {"FY2026": "1234567.895", "FY2025": "-98765.435"}, "tax_rate": "0.19"}]})
    out = tmp_path / "rand.xlsx"
    write_workbook(r, {}, out, mpm=mpm)
    res = verify_workbook(out, 1e-9)
    assert res.ok, res.failures
