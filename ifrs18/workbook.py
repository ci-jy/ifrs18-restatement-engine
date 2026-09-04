"""Write the restatement as an Excel workbook with live formulas, and verify it
by recalculating in headless LibreOffice.

Sheets: Original (as filed), Mapping (category, rule, paragraph and review flag
per line; the Category column is the reviewer's input), Restated (IFRS 18
layout; totals use SUMIF on Mapping so changing a category recalculates
everything), MPM (reconciliation with tax and NCI effects), Checks
(conservation proofs). A hidden ``_expected`` sheet stores the values Python
computed for every formula cell that matters, which ``verify_workbook``
compares after recalculation.
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .mpm import MPMDefinition, reconcile
from .restate import SUBTOTALS, Restatement
from .rules import CATEGORIES, CATEGORY_LABELS, Classification
from .statement import ITEM, PROFIT

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
SUB_FILL = PatternFill("solid", fgColor="E2EFDA")
NUM = "#,##0.00;(#,##0.00);-"


def _f(d: Decimal) -> float | int:
    return int(d) if d == d.to_integral_value() else float(d)


def _header(ws, values):
    ws.append(values)
    for c in ws[ws.max_row]:
        c.font, c.fill = BOLD, HEAD_FILL


def write_workbook(r: Restatement, classes: dict[str, Classification], path: str | Path,
                   mpm: MPMDefinition | None = None, entity_type: str = "general") -> dict[tuple[str, str], float]:
    stmt = r.statement
    periods = stmt.periods
    pcols = [get_column_letter(5 + i) for i in range(len(periods))]  # Original: E, F
    mcols = [get_column_letter(9 + i) for i in range(len(periods))]  # Mapping: I, J
    expected: dict[tuple[str, str], float] = {}
    wb = Workbook()

    # Original
    ws = wb.active
    ws.title = "Original"
    ws.append([f"{stmt.entity} - income statement as reported ({stmt.currency or 'reporting currency'}; income +, expense -)"])
    ws["A1"].font = BOLD
    ws.append([f"Source: {stmt.source}"])
    _header(ws, ["Line", "Label", "XBRL element", "Kind", *periods])
    orig_row = {}
    for ln in stmt.lines:
        ws.append([ln.line_id, ln.label, ln.tag, ln.kind, *[_f(ln.amounts.get(p, Decimal(0))) for p in periods]])
        orig_row[ln.line_id] = ws.max_row
        if ln.kind != ITEM:
            for c in ws[ws.max_row]:
                c.font = BOLD
    first, last = 4, ws.max_row
    profit_line = stmt.profit_line()

    # Mapping
    wm = wb.create_sheet("Mapping")
    _header(wm, ["Line", "Label", "Category", "Rule", "IFRS 18 ref", "Review flag", "Alternatives", "Reason", *periods])
    map_row = {}
    for ln in stmt.items():
        c = classes.get(ln.line_id)
        row = wm.max_row + 1
        wm.append([ln.line_id, ln.label, r.mapping[ln.line_id], c.rule_id if c else "manual", c.paragraph if c else "",
                   "REVIEW" if c and c.flagged else "", ", ".join(c.alternatives) if c else "", c.reason if c else "",
                   *[f"=Original!{pc}{orig_row[ln.line_id]}" for pc in pcols]])
        map_row[ln.line_id] = row
        if c and c.flagged:
            for cell in wm[row]:
                cell.fill = FLAG_FILL
    mfirst, mlast = 2, max(wm.max_row, 2)
    dv = DataValidation(type="list", formula1='"' + ",".join(CATEGORIES) + '"', allow_blank=False)
    wm.add_data_validation(dv)
    dv.add(f"C{mfirst}:C{mlast}")

    # Restated
    wr = wb.create_sheet("Restated")
    wr.append([f"{stmt.entity} - restated under IFRS 18 (entity type: {entity_type})"])
    wr["A1"].font = BOLD
    _header(wr, ["Line", "Label", *periods])
    cat_total_cell: dict[str, int] = {}
    sub_row: dict[str, int] = {}
    sub_after = {cats[-1]: (key, label) for key, label, cats in SUBTOTALS}
    for cat in CATEGORIES:
        wr.append(["", CATEGORY_LABELS[cat]])
        wr.cell(wr.max_row, 2).font = BOLD
        for ln in stmt.items():
            if r.mapping[ln.line_id] == cat:
                wr.append([ln.line_id, "  " + ln.label, *[f"=Mapping!{mc}{map_row[ln.line_id]}" for mc in mcols]])
        row = wr.max_row + 1
        wr.append(["", f"Total {CATEGORY_LABELS[cat].lower()}",
                   *[f'=ROUND(SUMIF(Mapping!$C${mfirst}:$C${mlast},"{cat}",Mapping!{mc}${mfirst}:{mc}${mlast}),2)' for mc in mcols]])
        wr.cell(row, 2).font = BOLD
        cat_total_cell[cat] = row
        for i, p in enumerate(periods):
            expected[("Restated", f"{get_column_letter(3 + i)}{row}")] = _f(r.category_totals[p][cat])
        key, label = sub_after[cat]
        prev = list(sub_row.values())
        row = wr.max_row + 1
        cols = [get_column_letter(3 + i) for i in range(len(periods))]
        wr.append(["", label, *[(f"=ROUND({c}{prev[-1]}+{c}{cat_total_cell[cat]},2)" if prev else f"={c}{cat_total_cell[cat]}") for c in cols]])
        for cell in wr[row]:
            cell.font, cell.fill = BOLD, SUB_FILL
        sub_row[key] = row
        for i, p in enumerate(periods):
            expected[("Restated", f"{cols[i]}{row}")] = _f(r.subtotals[p][key])
        wr.append([" "])
    if profit_line is not None:
        row = wr.max_row + 1
        wr.append(["", "Profit or loss as originally reported", *[f"=Original!{pc}{orig_row[profit_line.line_id]}" for pc in pcols]])
        wr.append(["", "Difference (must be zero)", *[f"=ROUND({get_column_letter(3 + i)}{sub_row['profit']}-{get_column_letter(3 + i)}{row},2)" for i in range(len(periods))]])
        for i in range(len(periods)):
            expected[("Restated", f"{get_column_letter(3 + i)}{row + 1}")] = 0
        for ln in stmt.lines:
            if ln.kind == "attribution":
                wr.append(["", "  " + ln.label, *[f"=Original!{pc}{orig_row[ln.line_id]}" for pc in pcols]])

    # Checks
    wc = wb.create_sheet("Checks")
    _header(wc, ["Check", *periods])
    item_sum = [f'ROUND(SUMIF(Original!$D${first}:$D${last},"item",Original!{pc}${first}:{pc}${last}),2)' for pc in pcols]
    rows = [("Category totals minus sum of original lines",
             [f"=ROUND(" + "+".join(f"Restated!{get_column_letter(3 + i)}{cat_total_cell[c]}" for c in CATEGORIES) + f"-{item_sum[i]},2)" for i in range(len(periods))]),
            ("Restated profit minus sum of original lines",
             [f"=ROUND(Restated!{get_column_letter(3 + i)}{sub_row['profit']}-{item_sum[i]},2)" for i in range(len(periods))])]
    if profit_line is not None:
        rows.append(("Restated profit minus reported profit",
                     [f"=ROUND(Restated!{get_column_letter(3 + i)}{sub_row['profit']}-Original!{pcols[i]}{orig_row[profit_line.line_id]},2)" for i in range(len(periods))]))
    rows.append(("Lines without a valid category", [f'=COUNTA(Mapping!$C${mfirst}:$C${mlast})-' + "-".join(
        f'COUNTIF(Mapping!$C${mfirst}:$C${mlast},"{c}")' for c in CATEGORIES) for _ in periods]))
    for name, formulas in rows:
        wc.append([name, *formulas])
        for i in range(len(periods)):
            expected[("Checks", f"{get_column_letter(2 + i)}{wc.max_row}")] = 0

    # MPM
    if mpm is not None:
        _write_mpm(wb, r, mpm, map_row, sub_row, orig_row, pcols, expected)

    # expected values for verification
    we = wb.create_sheet("_expected")
    we.append(["sheet", "cell", "value"])
    for (sheet, cell), v in expected.items():
        we.append([sheet, cell, v])
    we.sheet_state = "hidden"

    for sheet, widths in {"Original": [8, 60, 60, 12, 18, 18], "Mapping": [8, 50, 14, 30, 18, 12, 22, 60, 18, 18],
                          "Restated": [8, 60, 18, 18], "Checks": [50, 18, 18], "MPM": [50, 10, 10, 18, 16, 16, 18, 16, 16]}.items():
        if sheet in wb.sheetnames:
            for i, w in enumerate(widths, 1):
                wb[sheet].column_dimensions[get_column_letter(i)].width = w
    for wsx in wb.worksheets:
        for row in wsx.iter_rows():
            for cell in row:
                if cell.column >= 3 and isinstance(cell.value, (int, float)) or (isinstance(cell.value, str) and cell.value.startswith("=")):
                    cell.number_format = NUM
    for sheet in ("Mapping",):
        for row in wb[sheet].iter_rows(min_row=2):
            row[7].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
    return expected


def _write_mpm(wb, r, mpm, map_row, sub_row, orig_row, pcols, expected):
    stmt = r.statement
    periods = stmt.periods
    ws = wb.create_sheet("MPM")
    ws.append([f"{mpm.name} - reconciliation to {dict((k, l) for k, l, _ in SUBTOTALS)[mpm.subtotal].lower()} (IFRS 18.122-123)"])
    ws["A1"].font = BOLD
    if mpm.description:
        ws.append([mpm.description])
    hdr = ["Item", "Tax rate", "NCI share"]
    for p in periods:
        hdr += [f"{p} amount", f"{p} tax effect", f"{p} NCI effect"]
    _header(ws, hdr)
    # period column triplets start at D
    trip = [(get_column_letter(4 + 3 * i), get_column_letter(5 + 3 * i), get_column_letter(6 + 3 * i)) for i in range(len(periods))]
    recs = [reconcile(mpm, r, p) for p in periods]
    start = ws.max_row + 1
    rcols = [get_column_letter(3 + i) for i in range(len(periods))]
    row = [recs[0].subtotal_label + " (IFRS 18 subtotal)", "", ""]
    for i in range(len(periods)):
        row += [f"=ROUND(Restated!{rcols[i]}{sub_row[mpm.subtotal]},2)", "", ""]
    ws.append(row)
    for i, rec in enumerate(recs):
        expected[("MPM", f"{trip[i][0]}{start}")] = _f(rec.subtotal)
    adj_first = start + 1
    for j, adj in enumerate(mpm.adjustments):
        rr = ws.max_row + 1
        row = [adj.label, _f(adj.tax_rate), _f(adj.nci_share)]
        for i, p in enumerate(periods):
            a, t, n = trip[i]
            amt = f"=ROUND(-Original!{pcols[i]}{orig_row[adj.line]},2)" if adj.line else f"=ROUND({adj.amounts.get(p, Decimal(0))},2)"
            row += [amt, f"=ROUND(-{a}{rr}*$B{rr},2)", f"=ROUND(-({a}{rr}+{t}{rr})*$C{rr},2)"]
        ws.append(row)
        for i, rec in enumerate(recs):
            a, t, n = trip[i]
            expected[("MPM", f"{a}{rr}")] = _f(rec.rows[j].amount)
            expected[("MPM", f"{t}{rr}")] = _f(rec.rows[j].tax_effect)
            expected[("MPM", f"{n}{rr}")] = _f(rec.rows[j].nci_effect)
    adj_last = ws.max_row
    tr = ws.max_row + 1
    row = [mpm.name, "", ""]
    for i in range(len(periods)):
        a, t, n = trip[i]
        rng = lambda col: f"{col}{adj_first}:{col}{adj_last}" if adj_last >= adj_first else "0"
        row += [f"=ROUND({a}{start}+SUM({rng(a)}),2)" if adj_last >= adj_first else f"={a}{start}",
                f"=ROUND(SUM({rng(t)}),2)" if adj_last >= adj_first else 0, f"=ROUND(SUM({rng(n)}),2)" if adj_last >= adj_first else 0]
    ws.append(row)
    for cell in ws[tr]:
        cell.font, cell.fill = BOLD, SUB_FILL
    for i, rec in enumerate(recs):
        a, t, n = trip[i]
        expected[("MPM", f"{a}{tr}")] = _f(rec.mpm)
        expected[("MPM", f"{t}{tr}")] = _f(rec.total_tax)
        expected[("MPM", f"{n}{tr}")] = _f(rec.total_nci)
    ws.append([" "])
    ws.append(["Tax effect = -(adjustment x tax rate); NCI effect = -(adjustment + tax effect) x NCI share; all rounded to the cent."])


@dataclass
class VerifyResult:
    checked: int
    max_abs_diff: float
    failures: list[tuple[str, str, float, object]]

    @property
    def ok(self) -> bool:
        return not self.failures


def find_soffice() -> str:
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    raise FileNotFoundError("LibreOffice (soffice) not found; install it, e.g. apt-get install libreoffice-calc")


def recalculate(path: str | Path, outdir: str | Path) -> Path:
    """Open the workbook in headless LibreOffice, recalculate and save a copy."""
    path, outdir = Path(path).resolve(), Path(outdir)
    profile = Path(outdir) / "lo-profile"
    cmd = [find_soffice(), f"-env:UserInstallation=file://{profile}", "--headless", "--norestore",
           "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", str(outdir), str(path)]
    subprocess.run(cmd, check=True, capture_output=True, timeout=180)
    out = outdir / path.name
    if not out.exists():
        raise RuntimeError(f"LibreOffice did not write {out}")
    return out


def verify_workbook(path: str | Path, tolerance: float = 1e-9) -> VerifyResult:
    with tempfile.TemporaryDirectory(dir=os.environ.get("TMPDIR")) as tmp:
        out = recalculate(path, tmp)
        wb = load_workbook(out, data_only=True)
        exp = load_workbook(path, read_only=True)["_expected"]
        failures, checked, worst = [], 0, 0.0
        for sheet, cell, value in exp.iter_rows(min_row=2, values_only=True):
            got = wb[sheet][cell].value
            checked += 1
            if not isinstance(got, (int, float)):
                failures.append((sheet, cell, float(value), got))
                continue
            diff = abs(float(got) - float(value))
            worst = max(worst, diff)
            if diff > tolerance:
                failures.append((sheet, cell, float(value), got))
        return VerifyResult(checked, worst, failures)
